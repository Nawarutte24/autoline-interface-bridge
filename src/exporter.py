import openpyxl
from io import BytesIO
import os
import copy

def generate_output_excel(template_path, rows_to_write):
    """
    Copies Output Template (rows 1-6) and writes generated records starting from row 7.
    Returns: BytesIO object of the completed workbook.
    """
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    
    # Clear any sample data starting from row 7 downwards
    if ws.max_row >= 7:
        ws.delete_rows(7, ws.max_row - 6 + 10)
        
    current_row = 7
    for rec in rows_to_write:
        rec_type = rec.get("type")
        if rec_type == "BLANK":
            # Leave row empty
            current_row += 1
            continue
            
        for col_letter, val in rec.items():
            if col_letter == "type":
                continue
            if val is not None:
                col_idx = openpyxl.utils.column_index_from_string(col_letter)
                cell = ws.cell(row=current_row, column=col_idx)
                cell.value = val
                
                # Format dates
                if hasattr(val, "strftime"):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(val, (int, float)) and col_letter in ["I", "J", "M", "N"]:
                    cell.number_format = "#,##0.00"
                    
        current_row += 1
        
    output_stream = BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream
